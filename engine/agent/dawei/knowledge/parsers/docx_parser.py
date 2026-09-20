# Copyright (c) 2025 格律至微
# SPDX-License-Identifier: AGPL-3.0-only

"""DOCX parser via python-docx.

DOCX is a flow layout format with no inherent page boundaries, so page_count
is left None and page_offsets empty — the chunker's page annotation becomes a
no-op for DOCX (entity provenance will simply lack a page number).
"""

import logging
from pathlib import Path

from dawei.knowledge.models import Document, DocumentMetadata, DocumentType
from dawei.knowledge.parsers.base import BaseParser

logger = logging.getLogger(__name__)


class DocxParser(BaseParser):
    """Parser for .docx files (paragraphs + table cells)."""

    async def parse(self, file_path: str | Path) -> Document:
        file_path = Path(file_path)
        file_size = file_path.stat().st_size
        sha256_hash = self._calculate_sha256(file_path)

        # Lazy import — keeps parser module load cheap and avoids a hard
        # dependency at import time.
        from docx import Document as DocxDocument

        d = DocxDocument(str(file_path))

        blocks: list[str] = []
        for p in d.paragraphs:
            if p.text and p.text.strip():
                blocks.append(p.text.strip())
        # Flatten tables into the text stream
        for tbl in d.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    txt = cell.text.strip()
                    if txt:
                        blocks.append(txt)

        content = "\n\n".join(blocks)

        # Core properties (best-effort — not all docs carry them)
        title = None
        author = None
        try:
            cp = d.core_properties
            title = cp.title or None
            author = cp.author or None
        except Exception:
            pass

        if not content.strip():
            logger.warning(f"DOCX yielded no extractable text: {file_path.name}")

        metadata = DocumentMetadata(
            file_path=str(file_path),
            file_name=file_path.name,
            file_size=file_size,
            file_type=DocumentType.DOCX,
            sha256=sha256_hash,
            title=title,
            author=author,
        )
        return Document(id=sha256_hash, metadata=metadata, content=content)

    def supports_file_type(self, file_path: str | Path) -> bool:
        return Path(file_path).suffix.lower() == ".docx"
