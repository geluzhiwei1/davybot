# Copyright (c) 2025 格律至微
# SPDX-License-Identifier: AGPL-3.0-only

"""
IP 模块文档导出生成器。

支持四种格式：
  - DOCX (python-docx)   — 交底书 / 草案 / 商标材料 / 警告函
  - PDF  (reportlab)     — 草案 / 资产报告
  - XLSX (openpyxl)      — 商标商品清单 / 资产报告
  - ZIP  (zipfile)       — 侵权证据包 / 商标材料包

所有公共 generate_* 函数返回 bytes，供 FastAPI Response 直接返回。
数据结构来自 ip_routes.py 的 _create_mock_* 工厂及 Agent 真实输出。
"""

from __future__ import annotations

import io
import json
import logging
import zipfile
from datetime import datetime
from typing import Any

logger = logging.getLogger(__name__)

# ── 中文字体（reportlab 内置 CID，无需系统字体文件）───────────────

_PDF_FONT_REGISTERED = False


def _ensure_pdf_font() -> str:
    """注册 reportlab 中文 CID 字体，返回字体名。失败则回退 Helvetica。"""
    global _PDF_FONT_REGISTERED
    font_name = "Helvetica"
    if _PDF_FONT_REGISTERED:
        return "STSong-Light" if _PDF_FONT_REGISTERED is True else font_name
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont

        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        _PDF_FONT_REGISTERED = True  # type: ignore[assignment]
        return "STSong-Light"
    except Exception as e:
        logger.warning(f"[IP-EXPORT] Failed to register CJK font: {e}")
        _PDF_FONT_REGISTERED = "failed"  # type: ignore[assignment]
        return font_name


# ── 通用辅助 ─────────────────────────────────────────────────────


def _safe_get(d: Any, *keys: str, default: Any = None) -> Any:
    """安全嵌套取值。"""
    cur = d
    for k in keys:
        if not isinstance(cur, dict):
            return default
        cur = cur.get(k)
        if cur is None:
            return default
    return cur


def _ts() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _bytes_io_to_bytes(buf: io.BytesIO) -> bytes:
    return buf.getvalue()


# ============================================================
# DOCX 生成（python-docx）
# ============================================================


def _new_docx(title: str) -> tuple[Any, io.BytesIO]:
    """创建 docx 文档，写入标题，返回 (document, buffer)。"""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    # 标题
    h = doc.add_heading(title, level=0)
    for run in h.runs:
        run.font.size = Pt(20)
    doc.add_paragraph(f"生成时间：{_ts()}")
    doc.add_paragraph("")
    buf = io.BytesIO()
    return doc, buf


def _finalize_docx(doc: Any, buf: io.BytesIO) -> bytes:
    doc.save(buf)
    return _bytes_io_to_bytes(buf)


def generate_disclosure_docx(result: dict[str, Any]) -> bytes:
    """交底书 DOCX：按章节生成。"""
    title = result.get("title", "技术交底书")
    doc, buf = _new_docx(title)
    sections = result.get("sections", {})
    if isinstance(sections, dict):
        for _key, sec in sections.items():
            if isinstance(sec, dict):
                doc.add_heading(sec.get("title", "章节"), level=1)
                doc.add_paragraph(sec.get("content", ""))
            elif isinstance(sec, str):
                doc.add_paragraph(sec)
    doc.add_paragraph("")
    doc.add_paragraph("— 由 DavyBot 知识产权模块生成 —")
    return _finalize_docx(doc, buf)


def generate_draft_docx(result: dict[str, Any]) -> bytes:
    """专利申请文件 DOCX：说明书 + 权利要求。"""
    title = result.get("title", "专利申请文件")
    doc, buf = _new_docx(title)

    country = result.get("targetCountry", "")
    if country:
        doc.add_paragraph(f"目标国家/地区：{country}")
    doc.add_paragraph("")

    # 说明书
    spec = result.get("specification", {})
    if isinstance(spec, dict):
        doc.add_heading("说明书", level=1)
        for field in ("technicalField", "backgroundArt", "summary", "detailedDescription", "examples"):
            val = spec.get(field)
            if val:
                labels = {
                    "technicalField": "技术领域",
                    "backgroundArt": "背景技术",
                    "summary": "发明内容",
                    "detailedDescription": "具体实施方式",
                    "examples": "实施例",
                }
                doc.add_heading(labels.get(field, field), level=2)
                doc.add_paragraph(str(val))

    # 权利要求书
    claims = result.get("claims", [])
    if isinstance(claims, list) and claims:
        doc.add_heading("权利要求书", level=1)
        for claim in claims:
            if isinstance(claim, dict):
                num = claim.get("number", "?")
                content = claim.get("content", "")
                ctype = claim.get("type", "")
                tag = f"（{ctype}）" if ctype else ""
                doc.add_paragraph(f"{num}. {content}{tag}")

    # 合规检查
    cc = result.get("complianceCheck")
    if isinstance(cc, dict):
        doc.add_heading("合规检查", level=1)
        doc.add_paragraph(f"整体状态：{cc.get('overallStatus', 'N/A')}")
        issues = cc.get("issues", [])
        if isinstance(issues, list):
            for iss in issues:
                if isinstance(iss, dict):
                    doc.add_paragraph(
                        f"[{iss.get('severity', '?')}] {iss.get('message', '')}",
                        style="List Bullet",
                    )

    doc.add_paragraph("")
    doc.add_paragraph("— 由 DavyBot 知识产权模块生成 —")
    return _finalize_docx(doc, buf)


def generate_trademark_material_docx(result: dict[str, Any]) -> bytes:
    """商标申请材料 DOCX。"""
    name = result.get("name", "商标")
    doc, buf = _new_docx(f"商标注册申请材料 — {name}")
    doc.add_heading("推荐分类", level=1)
    classes = result.get("recommendedClasses", [])
    if isinstance(classes, list):
        for cls in classes:
            if isinstance(cls, dict):
                doc.add_heading(
                    f"第{cls.get('classNumber', '?')}类 — {cls.get('className', '')}",
                    level=2,
                )
                doc.add_paragraph(f"类型：{cls.get('type', '')}")
                doc.add_paragraph(f"理由：{cls.get('rationale', '')}")
                goods = cls.get("goods", [])
                if isinstance(goods, list):
                    doc.add_paragraph("商品/服务：")
                    for g in goods:
                        doc.add_paragraph(str(g), style="List Bullet")

    reg = result.get("registrability")
    if isinstance(reg, dict):
        doc.add_heading("可注册性评估", level=1)
        doc.add_paragraph(f"综合评分：{reg.get('overallScore', 'N/A')}")
        doc.add_paragraph(f"评级：{reg.get('overallRating', 'N/A')}")
        doc.add_paragraph(f"总结：{reg.get('summary', '')}")

    doc.add_paragraph("")
    doc.add_paragraph("— 由 DavyBot 知识产权模块生成 —")
    return _finalize_docx(doc, buf)


def generate_warning_letter_docx(result: dict[str, Any], alert_id: str) -> bytes:
    """侵权警告函 DOCX。"""
    # 找到对应 alert 的 result 项
    target = None
    for r in result.get("results", []):
        alert = r.get("alert") if isinstance(r, dict) else None
        if isinstance(alert, dict) and alert.get("id") == alert_id:
            target = r
            break
    if not target:
        target = {"infringerName": "未知", "infringerProduct": ""}

    infringer = target.get("infringerName", "未知")
    product = target.get("infringerProduct", "")
    patent = target.get("matchedMyPatentNumber", "")
    prob = target.get("infringementProbability", "")
    rate = target.get("featureMatchRate", "")

    doc, buf = _new_docx("专利侵权警告函")
    doc.add_paragraph("")
    doc.add_paragraph(f"致：{infringer}")
    doc.add_paragraph("")
    body = (
        f"我方经调查发现，贵方产品「{product}」涉嫌侵犯我方"
        f"{f'专利（专利号：{patent}）' if patent else '知识产权'}。"
        f"\n\n经技术比对，该产品的特征匹配率达 {rate}%，"
        f"侵权概率评估为「{prob}」。"
        f"\n\n现郑重要求贵方：\n"
        f"1. 立即停止生产、销售、许诺销售涉嫌侵权的产品；\n"
        f"2. 在收到本函后 15 个工作日内书面回复，说明相关情况；\n"
        f"3. 如未在上述期限内回复，我方将采取法律手段维护合法权益，"
        f"包括但不限于向法院提起诉讼。"
    )
    doc.add_paragraph(body)
    doc.add_paragraph("")
    doc.add_paragraph(f"日期：{_ts()}")
    doc.add_paragraph("（本警告函由 DavyBot 知识产权模块生成，正式发送前请由律师审核）")
    return _finalize_docx(doc, buf)


# ============================================================
# PDF 生成（reportlab）
# ============================================================


def _new_pdf_canvas(title: str) -> tuple[Any, io.BytesIO]:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4
    font = _ensure_pdf_font()
    c.setFont(font, 18)
    c.drawString(72, height - 72, title)
    c.setFont(font, 9)
    c.drawString(72, height - 90, f"生成时间：{_ts()}")
    c.setFont(font, 10)
    return c, buf


def _pdf_text_lines(
    c: Any, lines: list[tuple[str, int]], start_y: int, line_height: int = 16
) -> int:
    """在 PDF canvas 上逐行写文本，返回下一个 y 坐标。自动换页。"""
    from reportlab.lib.pagesizes import A4

    width, height = A4
    font = _ensure_pdf_font()
    y = start_y
    bottom = 72
    for text, size in lines:
        if y < bottom:
            c.showPage()
            c.setFont(font, size)
            y = height - 72
        c.setFont(font, size)
        c.drawString(72, y, str(text))
        y -= line_height
    return y


def generate_draft_pdf(result: dict[str, Any]) -> bytes:
    """专利申请文件 PDF。"""
    from reportlab.lib.pagesizes import A4

    title = result.get("title", "专利申请文件")
    c, buf = _new_pdf_canvas(title)
    width, height = A4
    y = height - 110

    font = _ensure_pdf_font()
    lines: list[tuple[str, int]] = []
    country = result.get("targetCountry", "")
    if country:
        lines.append((f"目标国家/地区：{country}", 10))
    lines.append(("", 10))

    spec = result.get("specification", {})
    if isinstance(spec, dict):
        labels = {
            "technicalField": "技术领域",
            "backgroundArt": "背景技术",
            "summary": "发明内容",
            "detailedDescription": "具体实施方式",
        }
        lines.append(("说明书", 13))
        for field, label in labels.items():
            val = spec.get(field)
            if val:
                lines.append((label, 11))
                # 长文本按 ~70 字截断避免溢出
                text = str(val)
                for i in range(0, len(text), 70):
                    lines.append((text[i : i + 70], 9))

    claims = result.get("claims", [])
    if isinstance(claims, list) and claims:
        lines.append(("", 10))
        lines.append(("权利要求书", 13))
        for claim in claims:
            if isinstance(claim, dict):
                lines.append(
                    (f"{claim.get('number', '?')}. {claim.get('content', '')}", 9)
                )

    y = _pdf_text_lines(c, lines, y, line_height=14)
    c.setFont(font, 8)
    c.drawString(72, 50, "— 由 DavyBot 知识产权模块生成 —")
    c.showPage()
    c.save()
    return _bytes_io_to_bytes(buf)


def generate_portfolio_pdf(result: dict[str, Any]) -> bytes:
    """IP 资产组合报告 PDF。"""
    from reportlab.lib.pagesizes import A4

    c, buf = _new_pdf_canvas("IP 资产组合报告")
    width, height = A4
    font = _ensure_pdf_font()
    y = height - 110

    lines: list[tuple[str, int]] = []
    kpis = result.get("kpis", {})
    if isinstance(kpis, dict):
        lines.append(("核心指标", 13))
        pat = kpis.get("patent", {})
        if isinstance(pat, dict):
            lines.append(
                (
                    f"专利：共 {pat.get('total', 0)} 件，"
                    f"已授权 {pat.get('granted', 0)}，"
                    f"审查中 {pat.get('pending', 0)}",
                    10,
                )
            )
        tm = kpis.get("trademark", {})
        if isinstance(tm, dict):
            lines.append(
                (
                    f"商标：共 {tm.get('total', 0)} 件，"
                    f"已注册 {tm.get('registered', 0)}，"
                    f"审查中 {tm.get('pending', 0)}",
                    10,
                )
            )
        lines.append((f"加权平均价值分：{kpis.get('weightedAvgValueScore', 'N/A')}", 10))

    health = result.get("healthScore")
    if health is not None:
        lines.append((f"组合健康度：{health}", 10))

    patents = result.get("patents", [])
    if isinstance(patents, list) and patents:
        lines.append(("", 10))
        lines.append((f"专利明细（{len(patents)} 件）", 13))
        for p in patents:
            if isinstance(p, dict):
                lines.append(
                    (
                        f"• {p.get('patentNumber', '')} {p.get('title', '')} "
                        f"[{p.get('status', '')}] 价值:{p.get('valueScore', '?')}",
                        9,
                    )
                )

    trademarks = result.get("trademarks", [])
    if isinstance(trademarks, list) and trademarks:
        lines.append(("", 10))
        lines.append((f"商标明细（{len(trademarks)} 件）", 13))
        for t in trademarks:
            if isinstance(t, dict):
                lines.append(
                    (
                        f"• {t.get('name', '')} [{t.get('status', '')}] "
                        f"类别:{t.get('classes', '')}",
                        9,
                    )
                )

    competitors = result.get("competitors", [])
    if isinstance(competitors, list) and competitors:
        lines.append(("", 10))
        lines.append(("竞品对比", 13))
        for comp in competitors:
            if isinstance(comp, dict):
                lines.append(
                    (
                        f"• {comp.get('name', '')}：专利 {comp.get('patentCount', 0)} 件",
                        9,
                    )
                )

    y = _pdf_text_lines(c, lines, y, line_height=14)
    c.setFont(font, 8)
    c.drawString(72, 50, "— 由 DavyBot 知识产权模块生成 —")
    c.showPage()
    c.save()
    return _bytes_io_to_bytes(buf)


# ============================================================
# Excel 生成（openpyxl）
# ============================================================


def generate_trademark_goods_xlsx(result: dict[str, Any]) -> bytes:
    """商标商品清单 Excel。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "商标商品清单"
    ws.append(["类别编号", "类别名称", "类型", "商品/服务项", "理由"])
    classes = result.get("recommendedClasses", [])
    if isinstance(classes, list):
        for cls in classes:
            if isinstance(cls, dict):
                goods = cls.get("goods", [])
                goods_str = "\n".join(str(g) for g in goods) if isinstance(goods, list) else str(goods)
                ws.append(
                    [
                        cls.get("classNumber", ""),
                        cls.get("className", ""),
                        cls.get("type", ""),
                        goods_str,
                        cls.get("rationale", ""),
                    ]
                )
    # 列宽
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 50
    ws.column_dimensions["E"].width = 40
    buf = io.BytesIO()
    wb.save(buf)
    return _bytes_io_to_bytes(buf)


def generate_portfolio_xlsx(result: dict[str, Any]) -> bytes:
    """IP 资产组合 Excel 报告。"""
    from openpyxl import Workbook

    wb = Workbook()

    # 专利表
    ws1 = wb.active
    ws1.title = "专利"
    ws1.append(
        [
            "专利号",
            "标题",
            "状态",
            "类型",
            "申请日",
            "授权日",
            "到期日",
            "发明人",
            "价值分",
            "同族大小",
            "被引次数",
        ]
    )
    for p in result.get("patents", []):
        if isinstance(p, dict):
            inventors = p.get("inventors", [])
            inv_str = ", ".join(inventors) if isinstance(inventors, list) else str(inventors)
            ws1.append(
                [
                    p.get("patentNumber", ""),
                    p.get("title", ""),
                    p.get("status", ""),
                    p.get("patentType", ""),
                    p.get("filingDate", ""),
                    p.get("grantDate", ""),
                    p.get("expiryDate", ""),
                    inv_str,
                    p.get("valueScore", ""),
                    p.get("familySize", ""),
                    p.get("citationCount", ""),
                ]
            )

    # 商标表
    ws2 = wb.create_sheet("商标")
    ws2.append(["商标名称", "注册号", "状态", "类别", "申请日", "注册日", "到期日"])
    for t in result.get("trademarks", []):
        if isinstance(t, dict):
            classes = t.get("classes", [])
            cls_str = ", ".join(str(c) for c in classes) if isinstance(classes, list) else str(classes)
            ws2.append(
                [
                    t.get("name", ""),
                    t.get("registrationNumber", ""),
                    t.get("status", ""),
                    cls_str,
                    t.get("filingDate", ""),
                    t.get("registrationDate", ""),
                    t.get("expiryDate", ""),
                ]
            )

    # KPI 表
    ws3 = wb.create_sheet("统计")
    kpis = result.get("kpis", {})
    if isinstance(kpis, dict):
        ws3.append(["指标", "值"])
        pat = kpis.get("patent", {})
        if isinstance(pat, dict):
            ws3.append(["专利总数", pat.get("total", "")])
            ws3.append(["已授权", pat.get("granted", "")])
            ws3.append(["审查中", pat.get("pending", "")])
        tm = kpis.get("trademark", {})
        if isinstance(tm, dict):
            ws3.append(["商标总数", tm.get("total", "")])
        ws3.append(["加权平均价值分", kpis.get("weightedAvgValueScore", "")])
    ws3.append(["组合健康度", result.get("healthScore", "")])

    buf = io.BytesIO()
    wb.save(buf)
    return _bytes_io_to_bytes(buf)


# ============================================================
# ZIP 生成（标准库 zipfile）
# ============================================================


def generate_evidence_zip(result: dict[str, Any], alert_id: str) -> bytes:
    """侵权证据包 ZIP：包含比对报告 + 警告函 + 结构化数据。"""
    target = None
    for r in result.get("results", []):
        alert = r.get("alert") if isinstance(r, dict) else None
        if isinstance(alert, dict) and alert.get("id") == alert_id:
            target = r
            break
    if target is None:
        target = {}

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        # 结构化证据数据
        zf.writestr(
            "evidence-data.json",
            json.dumps(target, ensure_ascii=False, indent=2, default=str),
        )
        # 文本比对报告
        infringer = target.get("infringerName", "未知")
        report_lines = [
            f"侵权证据比对报告",
            f"{'=' * 40}",
            f"侵权方：{infringer}",
            f"侵权产品：{target.get('infringerProduct', '')}",
            f"匹配专利：{target.get('matchedMyPatentNumber', '')}",
            f"侵权概率：{target.get('infringementProbability', '')}",
            f"特征匹配率：{target.get('featureMatchRate', '')}%",
            f"生成时间：{_ts()}",
            "",
            "权利要求比对明细：",
        ]
        for cm in target.get("claimMatches", []):
            if isinstance(cm, dict):
                report_lines.extend(
                    [
                        f"  - 权利要求 {cm.get('myClaimNumber', '?')}: "
                        f"{cm.get('matchLevel', '')}",
                        f"    我方：{cm.get('myClaimText', '')}",
                        f"    对方：{cm.get('theirFeature', '')}",
                        f"    分析：{cm.get('analysis', '')}",
                        "",
                    ]
                )
        zf.writestr("comparison-report.txt", "\n".join(report_lines))
        # 警告函
        warning_bytes = generate_warning_letter_docx(result, alert_id)
        zf.writestr(f"warning-letter-{alert_id}.docx", warning_bytes)
    return _bytes_io_to_bytes(buf)


def generate_trademark_package_zip(result: dict[str, Any]) -> bytes:
    """商标申请材料包 ZIP：DOCX 材料 + 商品清单 Excel + 结构化数据。"""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        # 商标材料 DOCX
        zf.writestr(
            "trademark-application.docx",
            generate_trademark_material_docx(result),
        )
        # 商品清单 Excel
        zf.writestr(
            "goods-list.xlsx",
            generate_trademark_goods_xlsx(result),
        )
        # 可注册性报告
        zf.writestr(
            "registrability-report.json",
            json.dumps(
                result.get("registrability", {}), ensure_ascii=False, indent=2, default=str
            ),
        )
    return _bytes_io_to_bytes(buf)
